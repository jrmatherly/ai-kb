"""Helpers for normalizing legacy or unusual Excel uploads."""

from __future__ import annotations

import io
import logging
import os
import subprocess
import sys
import tempfile
import time
from pathlib import Path
from typing import Literal, Optional

logger = logging.getLogger(__name__)

# Workbook kinds ConvertSpreadsheet accepts: DuckDB's read_xlsx reads only .xlsx.
CONVERTIBLE_SPREADSHEET_TYPES = frozenset({"xls", "xlsb"})

# openpyxl's write-only mode wrote about 150k cells/s on an M-series Mac, so
# 5M cells is roughly half a minute of work; larger sheets are refused rather
# than tying up a gRPC worker.
MAX_CONVERT_CELLS = 5_000_000

# calamine allocates a sheet's whole bounding box before Python sees a cell: a
# 5 KB file holding only A1 and XFD1048576 got the process OOM-killed. The
# conversion therefore runs in a child process under this address-space limit
# (enforced on Linux; macOS ignores RLIMIT_AS) and this timeout. A 4.8M-cell
# sheet peaked at 1.3 GB of address space in the docreader image; 1 GiB killed it.
CONVERT_MEMORY_LIMIT_BYTES = 2 * 1024 * 1024 * 1024
CONVERT_TIMEOUT_SECONDS = 100

_XLS_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_ZIP_MAGIC = b"PK\x03\x04"


def detect_excel_format(content: bytes) -> str | None:
    """Return pandas/excel format id: xlsx, xls, xlsb, ods, or None."""
    if not content:
        return None

    # Private pandas helper, absent from pandas' type stubs.
    from pandas.io.excel._base import inspect_excel_format  # pyrefly: ignore[missing-module-attribute]

    # pandas 3 accepts a path or a buffer here, not raw bytes.
    ext = inspect_excel_format(content_or_path=io.BytesIO(content))
    if ext in ("xlsx", "xls", "xlsb", "ods"):
        return ext
    if ext == "zip":
        return "xlsx"

    if content.startswith(_ZIP_MAGIC):
        return "xlsx"
    if len(content) >= len(_XLS_MAGIC) and content.startswith(_XLS_MAGIC):
        return "xls"
    return None


ExcelEngine = Literal["openpyxl", "xlrd", "odf", "calamine"]


def engine_for_format(ext: str | None) -> ExcelEngine:
    if ext == "xls":
        return "xlrd"
    # openpyxl reads only the XML workbook; .xlsb is binary records.
    if ext == "xlsb":
        return "calamine"
    if ext == "xlsx":
        return "openpyxl"
    if ext == "ods":
        return "odf"
    return "openpyxl"


def convert_excel_to_xlsx_bytes(content: bytes, suffix: str = ".xlsx") -> bytes | None:
    """Convert arbitrary spreadsheet bytes to XLSX using LibreOffice, if available."""
    soffice = find_soffice()
    if not soffice:
        return None

    max_attempts = 3
    for attempt in range(1, max_attempts + 1):
        with tempfile.TemporaryDirectory() as temp_dir, tempfile.TemporaryDirectory() as profile_dir:
            src = os.path.join(temp_dir, f"input{suffix}")
            with open(src, "wb") as handle:
                handle.write(content)

            user_installation = Path(profile_dir).as_uri()
            cmd = [
                soffice,
                "--headless",
                f"-env:UserInstallation={user_installation}",
                "--convert-to",
                "xlsx",
                "--outdir",
                temp_dir,
                src,
            ]
            try:
                result = subprocess.run(cmd, capture_output=True, timeout=120)
            except (OSError, subprocess.TimeoutExpired) as exc:
                logger.warning("LibreOffice convert failed to start: %s", exc)
                return None

            if result.returncode != 0:
                stderr = result.stderr.decode("utf-8", errors="ignore")
                logger.warning(
                    "LibreOffice convert failed (attempt %s/%s): %s",
                    attempt,
                    max_attempts,
                    stderr,
                )
                if attempt < max_attempts:
                    time.sleep(0.5 * attempt)
                    continue
                return None

            for name in os.listdir(temp_dir):
                if name.endswith(".xlsx"):
                    with open(os.path.join(temp_dir, name), "rb") as handle:
                        converted = handle.read()
                    logger.info(
                        "Converted spreadsheet via LibreOffice (%s -> xlsx, %d bytes)",
                        suffix,
                        len(converted),
                    )
                    return converted

            if attempt < max_attempts:
                time.sleep(0.5 * attempt)
    return None


def normalize_excel_bytes(content: bytes, file_type: str | None = None) -> bytes:
    """Return bytes readable by pandas, converting via LibreOffice when needed."""
    ext = detect_excel_format(content)
    if ext is not None:
        return content

    suffixes = []
    if file_type:
        suffixes.append(f".{file_type.lstrip('.')}")
    suffixes.extend([".xlsx", ".xls", ".et", ".csv"])
    seen: set[str] = set()
    for suffix in suffixes:
        if suffix in seen:
            continue
        seen.add(suffix)
        converted = convert_excel_to_xlsx_bytes(content, suffix=suffix)
        if converted and detect_excel_format(converted) is not None:
            return converted

    raise ValueError(
        "Unrecognized Excel file format; the file may be corrupt, encrypted, "
        "or not a spreadsheet"
    )


def find_soffice() -> Optional[str]:
    possible_paths = [
        "/usr/bin/soffice",
        "/usr/lib/libreoffice/program/soffice",
        "/opt/libreoffice25.2/program/soffice",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "C:\\Program Files\\LibreOffice\\program\\soffice.exe",
        "C:\\Program Files (x86)\\LibreOffice\\program\\soffice.exe",
    ]
    for path in possible_paths:
        if path and os.path.exists(path):
            return path

    result = subprocess.run(["which", "soffice"], capture_output=True, text=True)
    if result.returncode == 0 and result.stdout.strip():
        return result.stdout.strip()
    return None


def to_xlsx_bytes(content: bytes, max_cells: int = MAX_CONVERT_CELLS) -> bytes:
    """Rewrite any workbook calamine reads as .xlsx, in this process.

    Every sheet keeps its name, order and cell positions; numbers, booleans
    and dates stay typed. Formulas become their cached values. Call
    convert_spreadsheet_to_xlsx instead for untrusted input.
    """
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ERROR_CODES, ILLEGAL_CHARACTERS_RE
    from python_calamine import CalamineWorkbook, SheetTypeEnum

    book = CalamineWorkbook.from_filelike(io.BytesIO(content))
    out = Workbook(write_only=True)
    cells = 0
    for meta in book.sheets_metadata:
        if meta.typ != SheetTypeEnum.WorkSheet:
            continue
        sheet = book.get_sheet_by_name(meta.name)
        ws = out.create_sheet(meta.name)
        if sheet.start is None or sheet.end is None:
            continue
        # Written from A1, so the output grid reaches the last used cell.
        cells += (sheet.end[0] + 1) * (sheet.end[1] + 1)
        if cells > max_cells:
            raise ValueError(f"workbook has more than {max_cells} cells")
        # iter_rows starts at row 1 but at the first used column.
        left_pad = [None] * sheet.start[1]
        for row in sheet.iter_rows():
            values: list = list(left_pad)
            for value in row:
                if value == "":
                    value = None
                elif isinstance(value, str):
                    value = ILLEGAL_CHARACTERS_RE.sub("", value)
                    if value.startswith("=") or value in ERROR_CODES:
                        # openpyxl would store such text as a formula or an error cell.
                        cell = WriteOnlyCell(ws, value=value)
                        cell.data_type = "s"
                        value = cell
                values.append(value)
            ws.append(values)
    if not out.worksheets:
        raise ValueError("workbook has no worksheets")
    buf = io.BytesIO()
    out.save(buf)
    return buf.getvalue()


def convert_spreadsheet_to_xlsx(content: bytes, file_type: str, max_bytes: int) -> bytes:
    """Convert an untrusted .xls or .xlsb workbook to .xlsx bytes.

    Raises ValueError for unsupported, unreadable or oversized input.
    """
    file_type = file_type.lower().lstrip(".")
    if file_type not in CONVERTIBLE_SPREADSHEET_TYPES:
        raise ValueError(f"unsupported spreadsheet type {file_type!r}; expected xls or xlsb")
    if not content:
        raise ValueError("empty file")
    if len(content) > max_bytes:
        raise ValueError(f"file is {len(content)} bytes; the limit is {max_bytes}")
    try:
        # -I keeps the child from importing the docreader package (seconds of parser imports).
        result = subprocess.run(
            [sys.executable, "-I", __file__],
            input=content,
            capture_output=True,
            timeout=CONVERT_TIMEOUT_SECONDS,
        )
    except subprocess.TimeoutExpired as exc:
        raise ValueError(f"conversion took longer than {CONVERT_TIMEOUT_SECONDS}s") from exc
    if result.returncode < 0:
        raise ValueError(
            f"conversion was killed by signal {-result.returncode}; "
            "the workbook likely exceeds the memory limit"
        )
    if result.returncode != 0:
        lines = result.stderr.decode("utf-8", errors="replace").strip().splitlines()
        raise ValueError(lines[-1] if lines else f"conversion exited with {result.returncode}")
    # The response also carries a field tag and length prefix; keep it under the gRPC send limit.
    if len(result.stdout) + 16 > max_bytes:
        raise ValueError(f"converted workbook is {len(result.stdout)} bytes; the limit is {max_bytes}")
    return result.stdout


def _convert_child_main() -> int:
    try:
        import resource

        resource.setrlimit(
            resource.RLIMIT_AS, (CONVERT_MEMORY_LIMIT_BYTES, CONVERT_MEMORY_LIMIT_BYTES)
        )
    except (ImportError, ValueError, OSError):
        pass  # not enforceable here (macOS rejects or ignores RLIMIT_AS)
    try:
        # A container memory limit below the rlimit (Helm defaults to 512Mi) makes
        # the cgroup OOM killer act first; this makes it pick the child, not the server.
        with open("/proc/self/oom_score_adj", "w") as handle:
            handle.write("1000")
    except OSError:
        pass
    try:
        data = to_xlsx_bytes(sys.stdin.buffer.read())
    except MemoryError:
        print("workbook exceeds the conversion memory limit", file=sys.stderr)
        return 1
    except Exception as exc:  # noqa: BLE001 - the parent reports the message
        print(f"cannot read workbook: {exc}", file=sys.stderr)
        return 1
    sys.stdout.buffer.write(data)
    return 0


if __name__ == "__main__":
    sys.exit(_convert_child_main())

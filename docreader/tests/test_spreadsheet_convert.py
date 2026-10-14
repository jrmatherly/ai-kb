import datetime
import io
import unittest
from pathlib import Path

import openpyxl
from python_calamine import CalamineWorkbook

from docreader.parser.excel_convert import convert_spreadsheet_to_xlsx, to_xlsx_bytes

ROOT = Path(__file__).resolve().parents[2]
XLSB = Path(__file__).with_name("fixtures") / "pandas-test1.xlsb"
XLS = ROOT / "internal/infrastructure/docparser/anydoc/testdata/upstream/xls/sheet.xls"
LIMIT = 50 * 1024 * 1024


def as_read_back(value):
    """What openpyxl returns for a value calamine read: no empty strings, dates as datetimes."""
    if value == "":
        return None
    if type(value) is datetime.date:
        return datetime.datetime.combine(value, datetime.time())
    return value


def calamine_grid(content: bytes) -> dict:
    book = CalamineWorkbook.from_filelike(io.BytesIO(content))
    return {
        name: [[as_read_back(v) for v in row] for row in book.get_sheet_by_name(name).to_python()]
        for name in book.sheet_names
    }


def openpyxl_grid(content: bytes) -> dict:
    book = openpyxl.load_workbook(io.BytesIO(content))
    return {ws.title: [list(row) for row in ws.iter_rows(values_only=True)] for ws in book.worksheets}


class ConvertSpreadsheetTest(unittest.TestCase):
    def test_xlsb_keeps_sheets_and_values(self):
        content = XLSB.read_bytes()
        out = openpyxl_grid(convert_spreadsheet_to_xlsx(content, "xlsb", LIMIT))
        self.assertEqual(list(out), ["Sheet1", "Sheet2", "Sheet3", "Sheet4"])
        self.assertEqual(out["Sheet1"][0], [None, "A", "B", "C", "D"])
        self.assertEqual(out["Sheet1"][1][0], datetime.datetime(2000, 1, 3))
        self.assertEqual(out["Sheet1"][1][1], 0.980268513777)
        self.assertEqual(out["Sheet4"], [[None, "col1", "col2"], ["i1", "a", "x"], ["i2", "b", "y"]])
        self.assertEqual(out, calamine_grid(content))

    def test_xls_keeps_sheets_and_values(self):
        content = XLS.read_bytes()
        out = openpyxl_grid(convert_spreadsheet_to_xlsx(content, "XLS", LIMIT))
        self.assertEqual(list(out), ["Values", "Merged Grid"])
        self.assertEqual(out["Values"][1], ["Percent", 0.155, "fifteen and a half"])
        self.assertEqual(out["Values"][3], ["Thousands", 9876543, "grouped"])
        self.assertEqual(out["Merged Grid"], [["Span two", None, "Right"], [1, 2, 3]])

    def test_offset_cells_and_formula_like_text_keep_their_place(self):
        src = openpyxl.Workbook()
        ws = src.worksheets[0]
        ws.title = "Data"
        ws["C3"] = 7
        ws["D4"] = "=not a formula"
        ws["D4"].data_type = "s"
        ws["E4"] = "#N/A"
        ws["E4"].data_type = "s"
        buf = io.BytesIO()
        src.save(buf)
        xlsx = to_xlsx_bytes(buf.getvalue())
        out = openpyxl_grid(xlsx)
        self.assertEqual(out["Data"][2], [None, None, 7, None, None])
        self.assertEqual(out["Data"][3], [None, None, None, "=not a formula", "#N/A"])
        cells = openpyxl.load_workbook(io.BytesIO(xlsx))["Data"]
        self.assertEqual((cells["D4"].data_type, cells["E4"].data_type), ("s", "s"))

    def test_garbage_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "cannot read workbook"):
            convert_spreadsheet_to_xlsx(b"not a workbook at all", "xls", LIMIT)

    def test_unsupported_type_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "unsupported spreadsheet type"):
            convert_spreadsheet_to_xlsx(XLSB.read_bytes(), "csv", LIMIT)

    def test_oversize_input_is_refused(self):
        content = XLSB.read_bytes()
        with self.assertRaisesRegex(ValueError, "the limit is 1024"):
            convert_spreadsheet_to_xlsx(content, "xlsb", 1024)

    def test_cell_cap_is_enforced(self):
        with self.assertRaisesRegex(ValueError, "more than 10 cells"):
            to_xlsx_bytes(XLSB.read_bytes(), max_cells=10)


if __name__ == "__main__":
    unittest.main()
